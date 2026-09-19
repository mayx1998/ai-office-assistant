import json, os, re, subprocess
from graph_state import OfficeState
from langchain_core.messages import ToolMessage, AIMessage
from config import cfg


def try_repair_args(raw: str):
    """尽力把残缺 JSON 修成 dict，修不了返回 None。
    针对 qwen-plus 常丢结尾闭合符号的问题，逐个尝试补全。"""
    candidates = [raw, raw + "}", raw + "]}", raw + '"]}']
    last_err = None
    for c in candidates:
        try:
            obj = json.loads(c)
            if isinstance(obj, dict):
                return obj
        except Exception as e:
            last_err = e
            continue
    print(f"             [try_repair_args] 全部候选失败，最后一个错误: {last_err}")
    print(f"             [try_repair_args] raw 前100字符: {raw[:100]!r}")
    return None


# ── 第 3 刀的两个辅助函数 ─────────────────────────
def extract_text(file_path: str) -> str:
    """提取文件文本：docx 用 python-docx（含表格单元格），其余走 officecli。
    与 evaluator.py 里的版本保持一致。"""
    if file_path.endswith(".docx"):
        from docx import Document
        doc = Document(file_path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    result = subprocess.run(["officecli", "view", file_path, "text"],
                            capture_output=True, encoding="utf-8", errors="replace")
    return result.stdout


def find_output_files(task_text: str):
    """从任务文本里找出产物文件：绝对路径直接用；裸文件名用沙箱目录补全。
    data.xlsx 是输入文件，排除。"""
    abs_paths = re.findall(r"[A-Za-z]:\\[^\s，。；：、\"'<>|*?]*?\.(?:docx|xlsx|pptx)", task_text)
    sandbox_dir = os.path.dirname(abs_paths[0]) if abs_paths else None
    bare_names = re.findall(r"[\w一-鿿-]+\.(?:docx|xlsx|pptx)", task_text)
    candidates, seen = [], set()
    for p in abs_paths:
        if p not in seen:
            seen.add(p)
            candidates.append(p)
    if sandbox_dir:
        for n in bare_names:
            p = os.path.join(sandbox_dir, n)
            if p not in seen:
                seen.add(p)
                candidates.append(p)
    return [p for p in candidates if os.path.basename(p) != "data.xlsx"]


# ── planner：把任务拆成步骤 ─────────────────────
def make_planner(llm, tool_names):
    def planner(state: OfficeState) -> dict:
        print(f"\n═══ [planner] 开始规划 ═══")
        prompt = f"""将下面的文档任务拆成 2-5 个可执行步骤，每步对应一次工具调用。
        任务：{state['task']}
        可用工具：officecli（文档操作）、query_contact（查通讯录）

        规则：
        - 只描述要做什么，不要预先撰写正文内容（正文由执行阶段现场生成）
        - 步骤中涉及的文件路径、格式前缀（如「结论：」）、人名必须从任务原文照抄，不得省略、不得写成占位符
        - 只输出步骤列表，每行一步。
        注意：计划不能超出工具的实际能力。联系人查询工具只有张三、李四两个人，且只返回邮箱（没有部门、职位等信息）。"""
        resp = llm.invoke(prompt)
        plan = [l.strip() for l in resp.content.splitlines() if l.strip()]
        print(f"[planner] 拆出 {len(plan)} 步: {plan}")
        return {"plan": plan}
    return planner


# ── executor：执行当前步骤（复用 Day 1 的循环，缩小作用域）──
def make_executor(llm, tools):
    tool_map = {t.name: t for t in tools}
    llm_with_tools = llm.bind_tools(tools)

    def executor(state: OfficeState) -> dict:
        plan_text = "\n".join(state["plan"]) if state["plan"] else state["task"]
        # 第 1 刀：executor 直接看原始任务（计划降级为参考）
        messages = [("user", f"""原始任务（其中的文件路径、格式要求、前缀字样必须逐字满足）：{state['task']}

        参考计划：
        {plan_text}

        注意：计划只是参考。若计划与原始任务有出入（路径、格式前缀、具体内容），一律以原始任务为准。禁止把计划中的占位符（如"[邮箱]""内容待填充"）当作最终内容写入文件。请实际调用工具完成。""")]
        # 第 2 刀③：上轮 reviewer 的诊断（哪里错 + 怎么修）直接喂给本轮
        if state.get("review_feedback"):
            messages.append(("user", state["review_feedback"]))
        last_error = ""

        # 第 4 刀：防死循环——跟踪"连续相同错误"
        last_error_sig = None
        same_error_count = 0

        def note_result(tool_name, obs_text):
            """每次工具返回后调用。同一错误连续出现 >=2 次时，
            返回一条警告消息（调用方负责 append 进 messages），否则返回 None。"""
            nonlocal last_error_sig, same_error_count
            text = str(obs_text)
            is_error = (text.lstrip().startswith(("Error", "ERROR"))
                        or "not matched" in text
                        or "Traceback" in text)
            if not is_error:
                last_error_sig = None
                same_error_count = 0
                return None
            sig = (tool_name, text[:80])
            if sig == last_error_sig:
                same_error_count += 1
            else:
                same_error_count = 1
                last_error_sig = sig
            if same_error_count >= 2:
                return (f"警告：{tool_name} 同样的调用已连续报错 {same_error_count} 次，"
                        f"错误内容：{text[:150]}。禁止原样重试，必须换命令或参数写法。"
                        f"提示：读 xlsx 用 view <file> text 看全表，或 get <file> /Sheet1/B2 取单格；"
                        f"不要用 query 的坐标过滤语法。写单元格用 set <file> /Sheet1/A1 --prop value=内容。")
            return None

        for i in range(cfg()["limits"]["max_steps"]):                       # 步数上限放宽，多步任务需要
            resp = llm_with_tools.invoke(messages)
            if getattr(resp, "invalid_tool_calls", None):
                for itc in resp.invalid_tool_calls:
                    raw = itc.get("args") or ""
                    name = itc.get("name")
                    print(f"  [executor] 残缺调用: name={name}")
                    print(f"             args原文: {raw[:200]}")
                    fixed = try_repair_args(raw)
                    if fixed is not None and name in tool_map:
                        print(f"             ✓ 闭合补全修复成功，照常执行")
                        try:
                            obs = tool_map[name].invoke(fixed)
                        except Exception as e:
                            last_error = str(e)
                            obs = f"ERROR: {e}"
                        print(f"             返回: {str(obs)[:150]}")
                        if itc.get("id"):
                            messages.append(ToolMessage(content=str(obs), tool_call_id=itc["id"]))
                        else:
                            messages.append(("user", f"工具 {name} 已执行（参数经自动修复），结果：{str(obs)[:500]}"))
                        warn = note_result(name, obs)          # ← 第 4 刀
                        if warn:
                            messages.append(("user", warn))
                    else:
                        print(f"             无法修复，要求模型重试")
                        messages.append(AIMessage(content="（上一次工具调用参数 JSON 不合法，已丢弃）"))
                        messages.append(("user", "你的工具调用参数少了结尾的 } 或 ]，不是合法 JSON。请重新发起调用。"))
                continue
            messages.append(resp)

            if not resp.tool_calls:
                break
            for tc in resp.tool_calls:
                print(f"  [executor] 调用 {tc['name']}: {str(tc['args'])[:120]}")  # ← 日志
                try:
                    obs = tool_map[tc["name"]].invoke(tc["args"])
                except Exception as e:
                    last_error = str(e)
                    obs = f"ERROR: {e}"
                print(f"             返回: {str(obs)[:150]}")                      # ← 日志
                messages.append(ToolMessage(content=str(obs), tool_call_id=tc["id"]))
                warn = note_result(tc["name"], obs)              # ← 第 4 刀
                if warn:
                    messages.append(("user", warn))

        return {"messages": messages, "last_tool_error": last_error}
    return executor


# ── reviewer：三层校验 + 结构化诊断 ────────────────
def make_reviewer(llm):
    def reviewer(state: OfficeState) -> dict:
        # 第 1 层：有工具报错 → 不通过（附修复方向）
        if state.get("last_tool_error"):
            err = state["last_tool_error"]
            print(f"[reviewer] 存在工具报错，不通过: {err[:100]}")
            return {"review_pass": False,
                    "review_feedback": f"上轮执行出现工具报错：{err[:200]}。请换用正确的命令/参数重试；读 xlsx 用 view <file> text 或 get <file> /Sheet1/B2，不要反复重试同一种写法。"}

        tool_msgs = [m for m in state["messages"] if isinstance(m, ToolMessage)]

        # 第 2 层：没有任何工具执行记录 → 不通过
        if not tool_msgs:
            print("[reviewer] 没有任何工具执行记录，不通过")
            return {"review_pass": False,
                    "review_feedback": "上轮没有执行任何工具调用。请实际调用工具完成任务，不要只输出文字。"}

        # 第 2.5 层：只有查询、没有任何写操作痕迹 → 直接不通过，不问 LLM
        MUTATION_SIGNS = ["Created", "Added", "Updated", "saved", "写入"]
        if not any(any(sig in str(m.content) for sig in MUTATION_SIGNS) for m in tool_msgs):
            print("[reviewer] 只有查询记录、没有任何文件写操作，不通过")
            return {"review_pass": False,
                    "review_feedback": "上轮只有查询操作、没有写文件。任务要求产出文件，请用 officecli 实际创建文件并写入内容。"}

        # 证据 = 最近的工具回执 + 【第 3 刀】产物文件的真实内容
        evidence = "\n---\n".join(str(m.content)[:500] for m in tool_msgs[-6:])
        for p in find_output_files(state["task"])[:3]:       # 最多看 3 个产物文件
            name = os.path.basename(p)
            if os.path.exists(p):
                try:
                    evidence += f"\n---\n产物文件 {name} 的实际内容：\n{extract_text(p)[:1500]}"
                except Exception as e:
                    evidence += f"\n---\n产物文件 {name} 读取失败：{e}"
            else:
                evidence += f"\n---\n⚠️ 任务要求产物 {name}，但该文件在预期位置不存在（可能写错了目录）"

        # 第 3 层：LLM 对照 rubric 判定【第 2 刀：输出结构化 JSON，NO 必须带诊断】
        print(f"[reviewer] 提交 LLM 判定的证据:\n{evidence[:600]}")
        resp = llm.invoke(
            f"""任务：{state['task']}

证据（工具执行回执 + 产物文件实际内容）：
{evidence}

判定标准（严格遵守）：
- 对照任务逐条验收：要求的每个文件、每项具体内容（数值、人名、标题、格式前缀）都要在证据中找到
- 产物文件的实际内容优先级高于工具回执；内容里缺任务要求的东西 → 不通过
- 占位符（如"[邮箱]""内容待填充"）不算已完成的内容
- 不要脑补证据里没有的东西

只返回 JSON，不要输出任何其他文字：
{{"pass": true 或 false, "failed": ["没满足的具体要求"], "fix": "下一步该调用什么工具、写入什么内容来修复"}}"""
        )
        verdict = try_repair_args(resp.content.strip())   # 复用修复函数，兜住 JSON 截断
        if verdict is None:
            print(f"[reviewer] 判定 JSON 无法解析，按不通过处理。原文: {resp.content[:150]!r}")
            return {"review_pass": False,
                    "review_feedback": "上轮验收无法确认产物达标。请对照原始任务逐条检查产物文件，补齐缺漏后重新保存。"}

        passed = bool(verdict.get("pass"))
        feedback = ""
        if not passed:
            failed = verdict.get("failed") or ["未知"]
            fix = verdict.get("fix") or "请对照原始任务逐项检查产物文件"
            feedback = (f"上轮验收未通过。未满足的要求：{failed}；修复方向：{fix}。"
                        f"只修这些问题，不要重做已完成的部分，不要修改输入文件（data.xlsx 只读）。")
            print(f"[reviewer] LLM 判定: 不通过，原因: {failed}")
        else:
            print(f"[reviewer] LLM 判定: 通过")
        return {"review_pass": passed, "review_feedback": feedback}
    return reviewer


# ── repairer：带错误上下文重新规划 ──────────────
def make_repairer(llm):
    def repairer(state: OfficeState) -> dict:
        # 优先用 reviewer 的结构化诊断，其次才是工具报错
        error = state.get("review_feedback") or state["last_tool_error"] or "（无明确错误，reviewer 认为产物不符合任务要求）"
        resp = llm.invoke(f"""任务：{state['task']}
    问题反馈：{error}
    原计划：{state['plan']}
    请给出修正后的下一步具体操作（只输出一步）。
    注意：不要重做已完成的部分，不要修改输入文件（data.xlsx 只读）。""")
        return {"plan": [resp.content.strip()],
                "retry_count": state["retry_count"] + 1,
                "last_tool_error": ""}
    return repairer